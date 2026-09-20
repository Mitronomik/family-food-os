#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

EXPECTED_BASE = "d8c76a64483e3d5814e702be33c12cbe2e144160"
EXPECTED_PR39_CHECKPOINT = "a42beb529d9e3f4d219908f37fdb2ed430229cee43c18d414ff817ed1ac81b97"
EXPECTED_SHARD_ROWS = (1550, 1550, 1550, 1529)
EXPECTED_TOTAL_ROWS = 6179
EXPECTED_ELIGIBLE_ROWS = 6177
EXPECTED_CANDIDATES = 68
EXPECTED_RELEVANT_ROWS = 991
EXPECTED_EXTERNAL_IDS = 96

SOURCE_HASHES = {
    "russian_normative_recipes_v22_5_row_nutrients_part1.xlsx": "72a70f31b6b59454a94d78b73bdf2d43119f04799b266773bb34917e9cb3961e",
    "russian_normative_recipes_v22_5_row_nutrients_part2.xlsx": "5603956719b07aeeb93d4b348d08796dc23fc9b2a4f5a6b828da2619a576f9ae",
    "russian_normative_recipes_v22_5_row_nutrients_part3.xlsx": "7678d809f5ea1c68e1521d8460f213412f7fab13074ea45d9ce856fc37d4c11e",
    "russian_normative_recipes_v22_5_row_nutrients_part4.xlsx": "06bd3be3bf5c7b923808df08285032ee27d823f970e7ce047e4c56a6c8d0eb56",
    "russian_normative_recipes_v22_5_manifest.xlsx": "0dbd9983e64d5966ab5a182b38f912928e4a64d15de866d8409ace452e61468c",
    "russian_normative_recipes_v22_13_mass_nutrients.xlsx": "5ea78ead82568f8aff019a4076215c6598675783cb81e0d1d5b4f913016de4cc",
    "russian_normative_recipes_v22_13_integrity_audit.xlsx": "4353aec58e2610e3a9b4d46b970888cd17e1aabcd6ebea2af6ed0c323865d47f",
    "russian_normative_recipes_v22_13_manifest.xlsx": "f57eb053f11f6230ad733824dbf4d3bba523be13d9418c7ab0932e58046aa9f0",
}

CANDIDATE_CLASSES = {"DIRECT_EXISTING_MAP_LEAD", "CATALOGUE_EXTENSION_LEAD"}
EXISTING_STATES = {"EXACT_EXISTING", "ALIAS_EXISTING"}
EXTENSION_STATES = {"NEW_FOOD_CANDIDATE", "FORM_SPLIT_CANDIDATE"}

RELATIONSHIP_FIELDS = [
    "RelationshipSource", "SourceRecipeID", "RecipeNo", "RecipeName", "Category", "Variant",
    "OriginalIngredientID", "ResolvedIngredientID", "CanonicalIngredient", "OriginalIngredient",
    "Amount_g", "AmountStatus", "ChoiceGroup", "Optional", "RelationshipSourceURL",
    "NutrientInputEligible",
]

DUPLICATE_KEY_FIELDS = [
    "RelationshipSource", "SourceRecipeID", "Variant", "OriginalIngredientID", "ResolvedIngredientID",
    "OriginalIngredient", "Amount_g", "AmountStatus", "ChoiceGroup", "Optional",
    "RelationshipSourceURL", "NutrientInputEligible",
]

BOUNDARY_RE = re.compile(r"гарнир|соус|без\s+гарнир|без\s+соус", re.IGNORECASE)


def die(msg: str) -> None:
    raise SystemExit(f"DC1 VALIDATION ERROR: {msg}")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def git_blob_sha1_bytes(data: bytes) -> str:
    h = hashlib.sha1()
    h.update(f"blob {len(data)}\0".encode())
    h.update(data)
    return h.hexdigest()


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("ё", "е").casefold().split())


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "YES" if value else "NO"
    return str(value).strip()


def decimal_value(value: object, *, field: str) -> Decimal:
    if value is None or clean(value) == "":
        die(f"missing numeric {field}")
    try:
        d = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        die(f"invalid Decimal for {field}: {value!r}: {exc}")
    if not d.is_finite():
        die(f"non-finite Decimal for {field}: {value!r}")
    return d


def decimal_text(d: Decimal) -> str:
    s = format(d, "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s or "0"


def boolish(value: object) -> bool:
    return clean(value).upper() in {"TRUE", "YES", "1"}


def load_csv_objects(paths: Iterable[Path]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(paths):
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                die(f"CSV has no header: {path}")
            rows.extend({k: (v or "") for k, v in row.items()} for row in reader)
    return rows


def xlsx_table(path: Path, sheet: str) -> list[dict[str, object]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        die(f"missing sheet {sheet!r} in {path.name}")
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        die(f"empty sheet {sheet!r} in {path.name}")
    names = [clean(x) for x in header]
    if any(not x for x in names):
        die(f"blank header in {path.name}:{sheet}")
    out: list[dict[str, object]] = []
    for values in it:
        if not any(v is not None and clean(v) != "" for v in values):
            continue
        out.append({names[i]: values[i] if i < len(values) else None for i in range(len(names))})
    return out


def verify_source_hash(path: Path) -> str:
    digest = sha256(path)
    expected = SOURCE_HASHES.get(path.name)
    if expected and digest != expected:
        die(f"source hash mismatch for {path.name}: expected {expected}, got {digest}")
    return digest


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, "") for k in fields})


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def ordered_join(values: Iterable[str]) -> str:
    return " | ".join(sorted({v for v in values if v}))


def map_state_bucket(mapping: dict[str, str]) -> str:
    state = mapping["map_state"]
    if state in EXISTING_STATES:
        return "existing"
    if state == "NEW_FOOD_CANDIDATE":
        return "new"
    if state == "FORM_SPLIT_CANDIDATE":
        return "form"
    if state == "COMPOSITE_OR_PROCESS_OUTPUT":
        return "composite"
    if state == "UNRESOLVED":
        return "unresolved"
    if state == "REJECT_TECHNICAL":
        return "rejected"
    die(f"unknown map_state {state}")


def authority_strategy(mapping: dict[str, str], semantic_status: str) -> tuple[str, str, str]:
    state = mapping["map_state"]
    tier = mapping["external_exactness_tier"]
    if state in EXISTING_STATES:
        return (
            "REUSE_ACCEPTED_IDENTITY_MAPPING",
            "CURRENT_PROFILE_SUITABILITY_NOT_REVALIDATED_IN_DC1",
            "CURRENT_REPOSITORY_TRUTH",
        )
    if semantic_status == "V22_5_V22_13_LABEL_DIFFERENCE_REVIEW_REQUIRED":
        return (
            "IDENTITY_SEMANTICS_REVIEW_THEN_SOURCE_VERIFICATION",
            "NOT_VERIFIED_DC1",
            "PENDING_SELECTED_SOURCE_REVIEW",
        )
    if state == "FORM_SPLIT_CANDIDATE":
        return (
            "FORM_DECISION_THEN_SOURCE_VERIFICATION",
            "NOT_VERIFIED_DC1",
            "PENDING_SELECTED_SOURCE_REVIEW",
        )
    if tier == "C_PROXY":
        return (
            "IDENTITY_DECISION_THEN_SOURCE_VERIFICATION",
            "NOT_VERIFIED_DC1",
            "PENDING_SELECTED_SOURCE_REVIEW",
        )
    return (
        "PIN_COMPATIBLE_OFFICIAL_SOURCE",
        "NOT_VERIFIED_DC1",
        "PENDING_SELECTED_SOURCE_REVIEW",
    )


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--row-shard", action="append", required=True, dest="row_shards")
    ap.add_argument("--v22-5-manifest", required=True)
    ap.add_argument("--v22-13-mass", required=True)
    ap.add_argument("--v22-13-audit", required=True)
    ap.add_argument("--v22-13-manifest", required=True)
    ap.add_argument("--mapping-dir", required=True)
    ap.add_argument("--nutrition-seed")
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    row_shards = [Path(p) for p in args.row_shards]
    manifest5 = Path(args.v22_5_manifest)
    mass13 = Path(args.v22_13_mass)
    audit13 = Path(args.v22_13_audit)
    manifest13 = Path(args.v22_13_manifest)
    mapping_dir = Path(args.mapping_dir)
    out = Path(args.output)
    nutrition_seed = Path(args.nutrition_seed) if args.nutrition_seed else None

    if len(row_shards) != 4:
        die(f"expected exactly 4 row shards, got {len(row_shards)}")
    row_shards = sorted(row_shards, key=lambda p: p.name)

    source_hashes = {}
    for p in [*row_shards, manifest5, mass13, audit13, manifest13]:
        if not p.exists():
            die(f"missing source file: {p}")
        source_hashes[p.name] = verify_source_hash(p)

    # Manifest completeness: do not let partial spreadsheet rendering masquerade as a complete shard.
    manifest_rows = xlsx_table(manifest5, "Manifest")
    manifest_by_file = {clean(r["File"]): r for r in manifest_rows}
    expected_manifest_counts = []
    for shard, expected in zip(row_shards, EXPECTED_SHARD_ROWS):
        mr = manifest_by_file.get(shard.name)
        if not mr:
            die(f"v22.5 manifest does not list {shard.name}")
        declared = int(mr["Data rows"])
        if declared != expected:
            die(f"manifest row count for {shard.name}: expected {expected}, got {declared}")
        expected_manifest_counts.append(declared)
    if sum(expected_manifest_counts) != EXPECTED_TOTAL_ROWS:
        die("v22.5 manifest total does not reconcile to 6179")

    # Canonical accepted mapping/candidate package. Relevant rows may be split across arbitrary shards.
    candidate_paths = sorted(mapping_dir.glob("recipe-candidates-part*.csv"))
    mapping_paths = sorted(mapping_dir.glob("ingredient-mapping-part*.csv"))
    if not candidate_paths or not mapping_paths:
        die(f"missing accepted PR39 mapping CSV shards under {mapping_dir}")
    all_candidates = load_csv_objects(candidate_paths)
    all_mappings = load_csv_objects(mapping_paths)
    lead_candidates = [r for r in all_candidates if r.get("mapping_class") in CANDIDATE_CLASSES]
    if len(lead_candidates) != EXPECTED_CANDIDATES:
        die(f"candidate universe mismatch: expected {EXPECTED_CANDIDATES}, got {len(lead_candidates)}")
    candidate_ids = {r["source_recipe_id"] for r in lead_candidates}
    if len(candidate_ids) != EXPECTED_CANDIDATES:
        die("duplicate source_recipe_id in accepted candidate universe")
    if any(r.get("external_raw_status") != "READY_RAW" for r in lead_candidates):
        die("candidate universe contains a non-READY_RAW row; investigate accepted PR39 package")
    candidate_by_id = {r["source_recipe_id"]: r for r in lead_candidates}
    mapping_by_id = {r["external_ingredient_id"]: r for r in all_mappings}
    if len(mapping_by_id) != len(all_mappings):
        die("unexplained duplicate external_ingredient_id in accepted mapping input")

    # Read full raw XLSX row layer directly. Never use text-rendered spreadsheet prefixes.
    all_relationship_rows: list[dict[str, object]] = []
    shard_counts = []
    for shard, expected in zip(row_shards, EXPECTED_SHARD_ROWS):
        rows = xlsx_table(shard, "v22_RowNutrients")
        for row in rows:
            row["__source_shard"] = shard.name
        shard_counts.append(len(rows))
        if len(rows) != expected:
            die(f"row-shard completeness failure for {shard.name}: expected {expected}, got {len(rows)}")
        all_relationship_rows.extend(rows)
    if len(all_relationship_rows) != EXPECTED_TOTAL_ROWS:
        die(f"row-shard total mismatch: expected {EXPECTED_TOTAL_ROWS}, got {len(all_relationship_rows)}")

    relevant_rows = [r for r in all_relationship_rows if clean(r["SourceRecipeID"]) in candidate_ids]
    if len(relevant_rows) != EXPECTED_RELEVANT_ROWS:
        die(f"candidate relationship-row count mismatch: expected {EXPECTED_RELEVANT_ROWS}, got {len(relevant_rows)}")
    present_ids = {clean(r["SourceRecipeID"]) for r in relevant_rows}
    missing_candidates = sorted(candidate_ids - present_ids)
    if missing_candidates:
        die(f"candidate(s) lost from source relationship rows: {missing_candidates}")

    # No required source relationship may disappear into a zero. Amount must be present and positive for these eligible rows.
    relevant_by_recipe: dict[str, list[dict[str, object]]] = defaultdict(list)
    for r in relevant_rows:
        rid = clean(r["SourceRecipeID"])
        iid = clean(r["ResolvedIngredientID"])
        if not iid:
            die(f"blank ResolvedIngredientID in candidate {rid}")
        if iid not in mapping_by_id:
            die(f"source identity {iid} in {rid} is absent from accepted PR39 mapping")
        amount = decimal_value(r["Amount_g"], field=f"{rid}/{iid}/Amount_g")
        if amount <= 0:
            die(f"non-positive candidate ingredient amount in {rid}/{iid}: {amount}")
        if not boolish(r["NutrientInputEligible"]):
            die(f"candidate source row unexpectedly not NutrientInputEligible: {rid}/{iid}")
        relevant_by_recipe[rid].append(r)

    # Exact duplicate relationship rows are not silently tolerated.
    seen = set()
    duplicate_rows = []
    for r in relevant_rows:
        key = tuple(clean(r.get(f)) for f in DUPLICATE_KEY_FIELDS)
        if key in seen:
            duplicate_rows.append(key)
        seen.add(key)
    if duplicate_rows:
        die(f"unexplained duplicate candidate relationship rows: {duplicate_rows[:3]}")

    # v22.13 coverage/identity reference and integrity evidence.
    coverage_rows = xlsx_table(mass13, "v22_RecipeCoverage")
    coverage_by_id = {clean(r["SourceRecipeID"]): r for r in coverage_rows}
    refs13 = xlsx_table(mass13, "v22_IngredientRefs")
    refs13_by_id = {clean(r["IngredientID"]): r for r in refs13}
    integrity = xlsx_table(audit13, "IntegrityAudit")
    manifest13_integrity = xlsx_table(manifest13, "Integrity")
    if not any(clean(r.get("Check")) == "Contribution rows" and clean(r.get("Status")) == "PASS" and int(r.get("After")) == EXPECTED_TOTAL_ROWS for r in integrity):
        die("v22.13 audit does not prove 6179 contribution rows preserved")
    if not any(clean(r.get("Check")) == "Contribution row count" and clean(r.get("Status")) == "PASS" and int(r.get("After")) == EXPECTED_TOTAL_ROWS for r in manifest13_integrity):
        die("v22.13 manifest integrity does not prove 6179 contribution rows")

    # Relevant PR39 mapping state must reconcile exactly to accepted per-recipe aggregates.
    aggregate_mismatches = []
    compatibility_rows = []
    candidate_output = []
    external_ids_used = set()
    label_diff_ids = set()

    for rid in sorted(candidate_ids, key=lambda s: (int(s.split("-")[-1]), s)):
        rs = relevant_by_recipe[rid]
        accepted = candidate_by_id[rid]
        cov = coverage_by_id.get(rid)
        if not cov:
            die(f"candidate {rid} absent from v22.13 RecipeCoverage")
        if int(cov["CalcRows"]) != len(rs):
            die(f"v22.5/v22.13 calculation-row mismatch for {rid}: v22.5={len(rs)}, v22.13 CalcRows={cov['CalcRows']}")

        variants = sorted({clean(r["Variant"]) for r in rs})
        if any(not v for v in variants):
            die(f"blank Variant in candidate {rid}")
        if int(cov["VariantBlocks"]) != len(variants):
            die(f"variant-count mismatch for {rid}: v22.5={len(variants)}, v22.13={cov['VariantBlocks']}")

        buckets = Counter()
        row_ids = []
        for r in rs:
            iid = clean(r["ResolvedIngredientID"])
            row_ids.append(iid)
            external_ids_used.add(iid)
            buckets[map_state_bucket(mapping_by_id[iid])] += 1
        expected_counts = {
            "existing": int(accepted["mapped_existing_rows"]),
            "new": int(accepted["new_food_rows"]),
            "form": int(accepted["form_split_rows"]),
            "composite": int(accepted["composite_process_rows"]),
            "unresolved": int(accepted["unresolved_rows"]),
        }
        got_counts = {k: buckets.get(k, 0) for k in expected_counts}
        if got_counts != expected_counts:
            aggregate_mismatches.append({"source_recipe_id": rid, "expected": expected_counts, "actual": got_counts})

        choices = sorted({clean(r["ChoiceGroup"]) for r in rs if clean(r["ChoiceGroup"])})
        optional_rows = [r for r in rs if clean(r["Optional"]).upper() == "YES"]
        boundary_reasons = []
        boundary_text = " | ".join([clean(rs[0]["RecipeName"]), *variants])
        if BOUNDARY_RE.search(boundary_text):
            boundary_reasons.append("RECIPE_OR_VARIANT_GARNISH_SAUCE_BOUNDARY")
        if any(normalize_text(c) in {"sauce", "garnish"} for c in choices):
            boundary_reasons.append("CHOICE_GROUP_GARNISH_SAUCE_BOUNDARY")

        simple = len(variants) == 1 and not choices and not optional_rows and not boundary_reasons
        structural_status = "SINGLE_VARIANT_NO_EXPLICIT_ALTERNATIVE" if simple else "MULTI_OR_ALTERNATIVE_REVIEW_REQUIRED"

        # Per-variant mandatory set (non-optional, non-choice), and all alternatives union.
        per_variant_required: dict[str, set[str]] = {}
        per_variant_all: dict[str, set[str]] = {}
        for v in variants:
            vrows = [r for r in rs if clean(r["Variant"]) == v]
            per_variant_all[v] = {clean(r["ResolvedIngredientID"]) for r in vrows}
            per_variant_required[v] = {
                clean(r["ResolvedIngredientID"])
                for r in vrows
                if clean(r["Optional"]).upper() != "YES" and not clean(r["ChoiceGroup"])
            }
        common_required = set.intersection(*per_variant_required.values()) if per_variant_required else set()
        all_union = set().union(*per_variant_all.values()) if per_variant_all else set()
        choice_ids = {clean(r["ResolvedIngredientID"]) for r in rs if clean(r["ChoiceGroup"])}
        optional_ids = {clean(r["ResolvedIngredientID"]) for r in rs if clean(r["Optional"]).upper() == "YES"}
        alternative_only = all_union - common_required
        selected_variant_required = next(iter(per_variant_required.values())) if simple else set()

        # v22.5 label vs accepted v22.13 reference label is a compatibility review signal, not an identity remap.
        semantic_review_ids = []
        hard_label_example = []
        for iid in sorted(all_union):
            ref = refs13_by_id.get(iid)
            if not ref:
                die(f"candidate identity {iid} absent from v22.13 IngredientRefs")
            if clean(ref["Name"]) != mapping_by_id[iid]["external_name"]:
                die(f"PR39/v22.13 name mismatch for {iid}: {mapping_by_id[iid]['external_name']!r} vs {clean(ref['Name'])!r}")
            old_labels = {normalize_text(r["CanonicalIngredient"]) for r in rs if clean(r["ResolvedIngredientID"]) == iid}
            old_labels |= {normalize_text(r["OriginalIngredient"]) for r in rs if clean(r["ResolvedIngredientID"]) == iid}
            new_label = normalize_text(ref["Name"])
            if new_label not in old_labels:
                semantic_review_ids.append(iid)
                label_diff_ids.add(iid)
                if iid == "ING-0069":
                    hard_label_example.append("ING-0069:v22.5=Шпик,v22.13=Жир кулинарный")

        relationship_rows_13 = int(cov["RelationshipRows"])
        compatibility_status = (
            "CALC_ROWS_MATCH_RELATIONSHIP_ROWS_MATCH"
            if relationship_rows_13 == len(rs)
            else "CALC_ROWS_MATCH_V22_13_HAS_ADDITIONAL_NONCALC_RELATIONSHIPS"
        )
        compatibility_rows.append({
            "source_recipe_id": rid,
            "recipe_name": clean(rs[0]["RecipeName"]),
            "v22_5_candidate_rows": len(rs),
            "v22_13_calc_rows": int(cov["CalcRows"]),
            "v22_13_relationship_rows": relationship_rows_13,
            "v22_13_extra_noncalc_relationship_rows": relationship_rows_13 - len(rs),
            "v22_5_variant_count": len(variants),
            "v22_13_variant_blocks": int(cov["VariantBlocks"]),
            "compatibility_status": compatibility_status,
            "semantic_label_review_ids": ordered_join(semantic_review_ids),
            "explicit_label_conflict_example": ordered_join(hard_label_example),
            "row_level_v22_13_relationship_equivalence_proven": "YES" if relationship_rows_13 == len(rs) and not semantic_review_ids else "NO",
        })

        existing_ids = sorted(i for i in all_union if mapping_by_id[i]["map_state"] in EXISTING_STATES)
        dc2_ids = sorted(i for i in all_union if mapping_by_id[i]["map_state"] in EXTENSION_STATES)
        if any(mapping_by_id[i]["map_state"] not in EXISTING_STATES | EXTENSION_STATES for i in all_union):
            die(f"candidate {rid} contains a mapping state outside DC1 lead contract")

        if simple and not dc2_ids:
            dc3_batch = "DC3-A_SIMPLE_NO_DC2_WRITE"
        elif simple:
            dc3_batch = "DC3-B_SIMPLE_AFTER_DC2"
        else:
            dc3_batch = "DC3-C_VARIANT_BOUNDARY_REVIEW"

        candidate_output.append({
            "source_recipe_id": rid,
            "recipe_name": clean(rs[0]["RecipeName"]),
            "category": clean(rs[0]["Category"]),
            "external_raw_status": accepted["external_raw_status"],
            "mapping_class": accepted["mapping_class"],
            "strict_raw_triage": accepted["strict_raw_triage"],
            "source_relationship_rows_v22_5": len(rs),
            "source_calc_rows_v22_13": int(cov["CalcRows"]),
            "source_relationship_rows_v22_13": relationship_rows_13,
            "relationship_compatibility_status": compatibility_status,
            "source_variant_count": len(variants),
            "source_variants": ordered_join(variants),
            "choice_groups": ordered_join(choices),
            "optional_row_count": len(optional_rows),
            "boundary_review_reasons": ordered_join(boundary_reasons),
            "variant_selection_status": structural_status,
            "single_variant_required_ids": ordered_join(selected_variant_required),
            "common_required_across_variants_ids": ordered_join(common_required),
            "choice_group_ingredient_ids": ordered_join(choice_ids),
            "optional_ingredient_ids": ordered_join(optional_ids),
            "alternative_or_variant_only_ids": ordered_join(alternative_only),
            "all_variant_union_ingredient_ids": ordered_join(all_union),
            "existing_mapping_ids": ordered_join(existing_ids),
            "dc2_required_ids": ordered_join(dc2_ids),
            "semantic_label_review_ids": ordered_join(semantic_review_ids),
            "production_ready": "NO",
            "proposed_dc3_batch": dc3_batch,
        })

    if aggregate_mismatches:
        die(f"v22.5 rows do not reproduce accepted PR39 aggregate mapping counts: {aggregate_mismatches[:2]}")
    if len(external_ids_used) != EXPECTED_EXTERNAL_IDS:
        die(f"external identity count mismatch: expected {EXPECTED_EXTERNAL_IDS}, got {len(external_ids_used)}")

    # Build one conservative demand row per accepted external identity. Do not name-deduplicate identities.
    demand_rows = []
    source_rows_by_iid: dict[str, list[dict[str, object]]] = defaultdict(list)
    for r in relevant_rows:
        source_rows_by_iid[clean(r["ResolvedIngredientID"])].append(r)

    nutrition_codes = set()
    if nutrition_seed:
        with nutrition_seed.open(encoding="utf-8-sig", newline="") as f:
            nutrition_codes = {r["canonical_code"] for r in csv.DictReader(f)}

    simple_ids = {r["source_recipe_id"] for r in candidate_output if r["variant_selection_status"] == "SINGLE_VARIANT_NO_EXPLICIT_ALTERNATIVE"}
    category_counts = Counter(r["category"] for r in candidate_output)

    for iid in sorted(external_ids_used):
        mapping = mapping_by_id[iid]
        rs = source_rows_by_iid[iid]
        ref = refs13_by_id.get(iid)
        if not ref:
            die(f"missing v22.13 IngredientRefs row for {iid}")
        old_canonical = sorted({clean(r["CanonicalIngredient"]) for r in rs if clean(r["CanonicalIngredient"])})
        old_original = sorted({clean(r["OriginalIngredient"]) for r in rs if clean(r["OriginalIngredient"])})
        new_name = clean(ref["Name"])
        # Flag the identity if any retained candidate relationship row carries a v22.5
        # label that does not match the accepted v22.13 identity label. One matching
        # usage must not hide another mismatching usage of the same external ID.
        new_norm = normalize_text(new_name)
        row_label_mismatch = False
        for source_row in rs:
            row_labels = {normalize_text(source_row["CanonicalIngredient"]), normalize_text(source_row["OriginalIngredient"])} - {""}
            if new_norm not in row_labels:
                row_label_mismatch = True
                break
        semantic_status = "V22_5_V22_13_LABEL_DIFFERENCE_REVIEW_REQUIRED" if row_label_mismatch else "LABEL_MATCH_OR_COMPATIBLE_TEXT"
        recipes = sorted({clean(r["SourceRecipeID"]) for r in rs}, key=lambda s: (int(s.split("-")[-1]), s))
        simple_recipes = sorted(set(recipes) & simple_ids, key=lambda s: (int(s.split("-")[-1]), s))
        input_mass = sum((decimal_value(r["Amount_g"], field=f"{iid}/Amount_g") for r in rs), Decimal(0))
        search_strategy, profile_status, rights_status = authority_strategy(mapping, semantic_status)
        food_code = mapping["familyfoodos_food_code"]
        if food_code and nutrition_seed:
            profile_presence = "PRESENT_IN_PROVIDED_CURRENT_SEED_EXTRACT_NOT_FORM_REVALIDATED" if food_code in nutrition_codes else "NOT_PRESENT_IN_PROVIDED_SEED_EXTRACT_REVIEW_REQUIRED"
        elif food_code:
            profile_presence = "NOT_CHECKED_BY_GENERATOR"
        else:
            profile_presence = "NOT_APPLICABLE_NO_ACCEPTED_FOOD_MAPPING"

        # Proposed DC2 grouping is triage only, never publication approval.
        if mapping["map_state"] in EXISTING_STATES:
            dc2_batch = "REUSE_NO_DC2_WRITE"
        elif semantic_status == "V22_5_V22_13_LABEL_DIFFERENCE_REVIEW_REQUIRED" or mapping["map_state"] == "FORM_SPLIT_CANDIDATE" or mapping["external_exactness_tier"] == "C_PROXY":
            dc2_batch = "DC2-C_IDENTITY_FORM_REVIEW"
        elif len(recipes) >= 4 or len(simple_recipes) >= 2:
            dc2_batch = "DC2-A_HIGH_IMPACT"
        else:
            dc2_batch = "DC2-B_STANDARD_EXTENSION"

        demand_rows.append({
            "external_ingredient_id": iid,
            "v22_13_external_name": mapping["external_name"],
            "v22_5_canonical_labels": ordered_join(old_canonical),
            "v22_5_original_labels": ordered_join(old_original),
            "external_exactness_tier": mapping["external_exactness_tier"],
            "map_state": mapping["map_state"],
            "familyfoodos_food_code": food_code,
            "mapping_confidence": mapping["mapping_confidence"],
            "candidate_recipe_count": len(recipes),
            "candidate_recipe_ids": ordered_join(recipes),
            "candidate_row_count": len(rs),
            "candidate_input_mass_g": decimal_text(input_mass),
            "simple_candidate_recipe_count": len(simple_recipes),
            "simple_candidate_recipe_ids": ordered_join(simple_recipes),
            "semantic_compatibility_status": semantic_status,
            "accepted_identity_mapping_status": "REUSE_ACCEPTED_MAPPING" if mapping["map_state"] in EXISTING_STATES else "NO_ACCEPTED_FAMILYFOODOS_IDENTITY_YET",
            "current_profile_presence_status": profile_presence,
            "profile_suitability_for_recipe_form": profile_status,
            "authority_source_candidate": "CURRENT_REPOSITORY_PROFILE_PROVENANCE" if mapping["map_state"] in EXISTING_STATES else "",
            "authority_source_verification_status": "NOT_REVALIDATED_FOR_THIS_RECIPE_SOURCE_FORM" if mapping["map_state"] in EXISTING_STATES else "NOT_SELECTED_OR_VERIFIED_DC1",
            "authority_search_strategy": "" if mapping["map_state"] in EXISTING_STATES else search_strategy,
            "rights_status": rights_status,
            "production_ready": "NO",
            "proposed_dc2_batch": dc2_batch,
        })

    if len(demand_rows) != EXPECTED_EXTERNAL_IDS:
        die(f"food-demand rows must remain one-per-external-identity: expected {EXPECTED_EXTERNAL_IDS}, got {len(demand_rows)}")

    # Batch partitions: no omissions, no overlaps.
    dc2_groups: dict[str, list[str]] = defaultdict(list)
    for r in demand_rows:
        dc2_groups[r["proposed_dc2_batch"]].append(r["external_ingredient_id"])
    dc3_groups: dict[str, list[str]] = defaultdict(list)
    for r in candidate_output:
        dc3_groups[r["proposed_dc3_batch"]].append(r["source_recipe_id"])

    def assert_partition(groups: dict[str, list[str]], universe: set[str], label: str) -> None:
        flattened = [x for vals in groups.values() for x in vals]
        if set(flattened) != universe:
            die(f"{label} partition does not cover universe exactly")
        dup = [k for k, n in Counter(flattened).items() if n != 1]
        if dup:
            die(f"{label} partition has overlaps/duplicates: {dup}")

    assert_partition(dc2_groups, external_ids_used, "DC2")
    assert_partition(dc3_groups, candidate_ids, "DC3")

    simple_candidates = [r for r in candidate_output if r["variant_selection_status"] == "SINGLE_VARIANT_NO_EXPLICIT_ALTERNATIVE"]
    review_candidates = [r for r in candidate_output if r["variant_selection_status"] != "SINGLE_VARIANT_NO_EXPLICIT_ALTERNATIVE"]
    if len(simple_candidates) != 26 or len(review_candidates) != 42:
        die(f"structural variant classification drift: expected 26/42, got {len(simple_candidates)}/{len(review_candidates)}")
    if any(r["proposed_dc3_batch"] != "DC3-C_VARIANT_BOUNDARY_REVIEW" for r in review_candidates):
        die("a candidate with unresolved alternatives was assigned an exact/simple DC3 status")

    existing_demands = [r for r in demand_rows if r["map_state"] in EXISTING_STATES]
    work_demands = [r for r in demand_rows if r["map_state"] in EXTENSION_STATES]
    if len(existing_demands) != 33 or len(work_demands) != 63:
        die(f"accepted mapping split drift: expected 33/63, got {len(existing_demands)}/{len(work_demands)}")

    compat_gap = [r for r in compatibility_rows if r["v22_13_extra_noncalc_relationship_rows"] > 0]
    exact_label_diffs = sorted(label_diff_ids)

    # Source relationship extract intentionally excludes nutrient values.
    relationship_output = []
    for r in sorted(relevant_rows, key=lambda x: (
        int(clean(x["SourceRecipeID"]).split("-")[-1]), clean(x["Variant"]), clean(x["ResolvedIngredientID"]), clean(x["OriginalIngredient"]), decimal_value(x["Amount_g"], field="sort amount")
    )):
        relationship_output.append({
            "source_shard": clean(r["__source_shard"]),
            "relationship_source": clean(r["RelationshipSource"]),
            "source_recipe_id": clean(r["SourceRecipeID"]),
            "recipe_no": clean(r["RecipeNo"]),
            "recipe_name": clean(r["RecipeName"]),
            "category": clean(r["Category"]),
            "variant": clean(r["Variant"]),
            "original_ingredient_id": clean(r["OriginalIngredientID"]),
            "resolved_ingredient_id": clean(r["ResolvedIngredientID"]),
            "v22_5_canonical_ingredient": clean(r["CanonicalIngredient"]),
            "original_ingredient": clean(r["OriginalIngredient"]),
            "amount_g": decimal_text(decimal_value(r["Amount_g"], field="relationship Amount_g")),
            "amount_status": clean(r["AmountStatus"]),
            "choice_group": clean(r["ChoiceGroup"]),
            "optional": clean(r["Optional"]),
            "relationship_source_url": clean(r["RelationshipSourceURL"]),
            "nutrient_input_eligible": clean(r["NutrientInputEligible"]),
        })

    # Assortment evidence: no auto-expansion.
    assortment = {
        "category_counts": dict(sorted(category_counts.items())),
        "assessment": "INITIAL_68_FUNNEL_SKEWED_REVIEW_REQUIRED",
        "observations": [
            "Candidate funnel is dominated by soups and potato/vegetable/mushroom dishes.",
            "Egg dishes are well represented relative to standalone meat/poultry families.",
            "This DC1 package preserves the 68-family funnel and does not expand it automatically; realistic weekly variety remains a DC3 curation blocker.",
        ],
    }

    summary = {
        "operation": "DATA-CORPUS-V1-DC1-RECOVERED",
        "repository_base": EXPECTED_BASE,
        "execution_issue": 67,
        "candidate_selection": {
            "rule": "All accepted PR39 DIRECT_EXISTING_MAP_LEAD + CATALOGUE_EXTENSION_LEAD recipes",
            "recipe_family_count": len(candidate_output),
            "ready_raw_external_status_count": len(candidate_output),
            "ready_raw_is_production_ready": False,
            "strict_raw_triage_yes": sum(r["strict_raw_triage"] == "YES" for r in candidate_output),
            "strict_raw_triage_no": sum(r["strict_raw_triage"] != "YES" for r in candidate_output),
            "source_variant_structure": {
                "one_source_variant": sum(r["source_variant_count"] == 1 for r in candidate_output),
                "multiple_source_variants": sum(r["source_variant_count"] > 1 for r in candidate_output),
                "single_variant_no_explicit_alternative": len(simple_candidates),
                "multi_or_alternative_review_required": len(review_candidates),
            },
            "source_relationship_rows_v22_5": len(relevant_rows),
            "v22_13_extra_noncalc_relationship_gap_recipe_count": len(compat_gap),
            "dc3_batch_counts": dict(sorted(Counter(r["proposed_dc3_batch"] for r in candidate_output).items())),
        },
        "food_demand": {
            "external_identity_count": len(external_ids_used),
            "deduplication_policy": "NO_NAME_BASED_DEDUP; one demand row per accepted external identity until equivalence is proven",
            "demand_row_count": len(demand_rows),
            "accepted_existing_identity_mapping_count": len(existing_demands),
            "dc2_required_identity_count": len(work_demands),
            "mapping_state_counts": dict(sorted(Counter(r["map_state"] for r in demand_rows).items())),
            "semantic_label_difference_review_count": len(exact_label_diffs),
            "semantic_label_difference_review_ids": exact_label_diffs,
            "dc2_batch_counts": dict(sorted(Counter(r["proposed_dc2_batch"] for r in demand_rows).items())),
        },
        "compatibility": {
            "v22_5_candidate_calc_rows_match_v22_13_calc_rows_for_all_68": True,
            "v22_5_rows_reproduce_pr39_mapping_aggregates_for_all_68": True,
            "candidate_recipes_with_v22_13_additional_noncalc_relationship_rows": len(compat_gap),
            "exact_v22_13_row_shards_available_to_this_recovery": False,
            "row_level_full_relationship_equivalence_claimed": False,
        },
        "assortment": assortment,
        "stop_condition": "DC1 evidence only. DC2/DC3 remain NOT STARTED and require separate reviewable production batches.",
    }

    batch_plan = {
        "dc2": {k: {"count": len(v), "external_ingredient_ids": sorted(v)} for k, v in sorted(dc2_groups.items())},
        "dc3": {k: {"count": len(v), "source_recipe_ids": sorted(v, key=lambda s: (int(s.split('-')[-1]), s))} for k, v in sorted(dc3_groups.items())},
        "rules": {
            "dc2": {
                "REUSE_NO_DC2_WRITE": "Accepted existing identity mapping; no DC2 write authorized by this package.",
                "DC2-A_HIGH_IMPACT": "Non-proxy/new identity with >=4 candidate families or >=2 structurally simple candidate families; source verification still required.",
                "DC2-B_STANDARD_EXTENSION": "Other non-proxy new identities; source verification still required.",
                "DC2-C_IDENTITY_FORM_REVIEW": "Form split, proxy, or v22.5/v22.13 label-difference item; identity/form semantics first.",
            },
            "dc3": {
                "DC3-A_SIMPLE_NO_DC2_WRITE": "One source variant, no choice/optional/boundary signal, and no DC2 identity dependency. Still not production-ready.",
                "DC3-B_SIMPLE_AFTER_DC2": "One source variant with no explicit alternative/boundary signal, but has DC2 identity dependency.",
                "DC3-C_VARIANT_BOUNDARY_REVIEW": "Multiple variants, ChoiceGroup, optional row, or garnish/sauce boundary review required before any exact publication branch.",
            },
        },
    }

    source_artifacts = {
        "repository": {
            "base_main": EXPECTED_BASE,
            "accepted_mapping_pr": 39,
            "accepted_mapping_checkpoint_sha256": EXPECTED_PR39_CHECKPOINT,
            "mapping_input_contract": [
                "data/curation/v22-13-map-a/recipe-candidates-part*.csv",
                "data/curation/v22-13-map-a/ingredient-mapping-part*.csv",
            ],
            "mapping_rows_loaded": {
                "candidate_rows_total_from_provided_shards": len(all_candidates),
                "ingredient_mapping_rows_total_from_provided_shards": len(all_mappings),
                "relevant_candidate_rows": len(lead_candidates),
                "relevant_external_identity_rows": len(external_ids_used),
            },
        },
        "local_source_files": [
            {
                "name": name,
                "sha256": digest,
                "use": (
                    "raw relationship/calculation-row extraction only; nutrient values not production authority"
                    if "v22_5_row_nutrients" in name
                    else "v22.5 shard-count/join-contract validation"
                    if name == "russian_normative_recipes_v22_5_manifest.xlsx"
                    else "v22.13 coverage/identity cross-check only; nutrient values not production authority"
                    if name == "russian_normative_recipes_v22_13_mass_nutrients.xlsx"
                    else "v22.13 integrity/semantic-repair cross-check only"
                    if name == "russian_normative_recipes_v22_13_integrity_audit.xlsx"
                    else "v22.13 checkpoint/integrity/file-inventory cross-check only"
                ),
            }
            for name, digest in sorted(source_hashes.items())
        ],
        "missing_source_artifacts": [
            "russian_normative_recipes_v22_13_row_nutrients_part1.xlsx",
            "russian_normative_recipes_v22_13_row_nutrients_part2.xlsx",
            "russian_normative_recipes_v22_13_row_nutrients_part3.xlsx",
            "russian_normative_recipes_v22_13_row_nutrients_part4.xlsx",
        ],
        "missing_artifact_effect": "Full row-level equality between v22.5 relationship rows and v22.13 relationship rows cannot be claimed. v22.5 candidate rows match v22.13 CalcRows and accepted PR39 aggregate mapping counts; extra non-calc v22.13 relationship rows remain an explicit compatibility blocker.",
        "production_nutrition_promotion": "NONE",
    }

    out.mkdir(parents=True, exist_ok=True)
    candidate_fields = [
        "source_recipe_id", "recipe_name", "category", "external_raw_status", "mapping_class", "strict_raw_triage",
        "source_relationship_rows_v22_5", "source_calc_rows_v22_13", "source_relationship_rows_v22_13",
        "relationship_compatibility_status", "source_variant_count", "source_variants", "choice_groups", "optional_row_count",
        "boundary_review_reasons", "variant_selection_status", "single_variant_required_ids",
        "common_required_across_variants_ids", "choice_group_ingredient_ids", "optional_ingredient_ids",
        "alternative_or_variant_only_ids", "all_variant_union_ingredient_ids", "existing_mapping_ids", "dc2_required_ids",
        "semantic_label_review_ids", "production_ready", "proposed_dc3_batch",
    ]
    demand_fields = [
        "external_ingredient_id", "v22_13_external_name", "v22_5_canonical_labels", "v22_5_original_labels",
        "external_exactness_tier", "map_state", "familyfoodos_food_code", "mapping_confidence", "candidate_recipe_count",
        "candidate_recipe_ids", "candidate_row_count", "candidate_input_mass_g", "simple_candidate_recipe_count",
        "simple_candidate_recipe_ids", "semantic_compatibility_status", "accepted_identity_mapping_status",
        "current_profile_presence_status", "profile_suitability_for_recipe_form", "authority_source_candidate",
        "authority_source_verification_status", "authority_search_strategy", "rights_status", "production_ready", "proposed_dc2_batch",
    ]
    compatibility_fields = [
        "source_recipe_id", "recipe_name", "v22_5_candidate_rows", "v22_13_calc_rows", "v22_13_relationship_rows",
        "v22_13_extra_noncalc_relationship_rows", "v22_5_variant_count", "v22_13_variant_blocks", "compatibility_status",
        "semantic_label_review_ids", "explicit_label_conflict_example", "row_level_v22_13_relationship_equivalence_proven",
    ]
    relationship_fields = [
        "source_shard", "relationship_source", "source_recipe_id", "recipe_no", "recipe_name", "category", "variant",
        "original_ingredient_id", "resolved_ingredient_id", "v22_5_canonical_ingredient", "original_ingredient",
        "amount_g", "amount_status", "choice_group", "optional", "relationship_source_url", "nutrient_input_eligible",
    ]

    write_csv(out / "candidate-recipes.csv", candidate_output, candidate_fields)
    write_csv(out / "food-demand.csv", demand_rows, demand_fields)
    write_csv(out / "compatibility.csv", compatibility_rows, compatibility_fields)
    relationship_part_names = []
    for part_no, shard in enumerate(row_shards, start=1):
        part_name = f"source-relationships-part{part_no}.csv"
        relationship_part_names.append(part_name)
        part_rows = [r for r in relationship_output if r["source_shard"] == shard.name]
        write_csv(out / part_name, part_rows, relationship_fields)
    write_json(out / "batch-plan.json", batch_plan)
    write_json(out / "source-artifacts.json", source_artifacts)
    write_json(out / "summary.json", summary)

    # README numbers are generated solely from summary/data.
    cat_lines = "\n".join(f"- {k}: **{v}**" for k, v in assortment["category_counts"].items())
    dc2_lines = "\n".join(f"- `{k}`: **{v}**" for k, v in summary["food_demand"]["dc2_batch_counts"].items())
    dc3_lines = "\n".join(f"- `{k}`: **{v}**" for k, v in summary["candidate_selection"]["dc3_batch_counts"].items())
    readme = f"""# DATA-CORPUS-V1 / DC1 — recovered source authority + coverage inventory

**Status:** evidence/curation only for Issue #67.
**Repository base:** `{EXPECTED_BASE}` (merged PR #69).
**Production publication:** **NO**.
**DC2 / DC3:** **NOT STARTED**.

## Recovery result

This package was rebuilt from raw XLSX bytes and accepted PR #39 mapping files after a completeness defect was found in the interrupted DC1 package.

The prior package used a text-rendered spreadsheet view as if it were the complete row shard. The raw v22.5 manifest requires four shards with **1550 / 1550 / 1550 / 1529 = 6179** rows. Raw XLSX parsing finds **{len(relevant_rows)}** candidate relationship/calculation rows for all **{len(candidate_output)}** candidates. The corrupted package had silently treated missing rendered rows as zero ingredient demand.

The preserved pre-rebuild package is under `recovery/pre-rebuild-b7bc19e8881ddc90/`.

## FACT — candidate universe

The candidate universe is unchanged: all accepted PR #39 `DIRECT_EXISTING_MAP_LEAD` + `CATALOGUE_EXTENSION_LEAD` rows.

- recipe families: **{len(candidate_output)}**;
- PR39 external status `READY_RAW`: **{len(candidate_output)}**; this is **not** production readiness;
- strict raw triage YES: **{summary['candidate_selection']['strict_raw_triage_yes']}**;
- one source `Variant`: **{summary['candidate_selection']['source_variant_structure']['one_source_variant']}**;
- multiple source `Variant` values: **{summary['candidate_selection']['source_variant_structure']['multiple_source_variants']}**;
- structurally simple single-variant families after ChoiceGroup/optional/boundary review: **{len(simple_candidates)}**;
- families requiring variant/choice/optional/boundary review: **{len(review_candidates)}**.

A single `Variant` value is not treated as an exact publication decision. Five one-variant families still require review because of explicit ChoiceGroup or garnish/sauce boundaries.

## FACT — source compatibility

For all 68 candidates:

- v22.5 candidate rows equal v22.13 `CalcRows`;
- v22.5 rows reproduce the accepted PR #39 mapping-state row aggregates;
- v22.5 variant count equals v22.13 `VariantBlocks`.

However **{len(compat_gap)}** candidates have more v22.13 `RelationshipRows` than v22.5 calculation rows. Exact v22.13 row-shard files were not available to this recovery environment. Therefore full row-level relationship equivalence is **not** claimed.

`compatibility.csv` carries this blocker explicitly.

## FACT — food/form demand

This recovery preserves external identity boundaries instead of deduplicating by label:

- external ingredient IDs used: **{len(demand_rows)}**;
- accepted existing/alias mappings: **{len(existing_demands)}**;
- identities requiring DC2-level identity/form/source work: **{len(work_demands)}**.

The earlier journal number 95 came from collapsing two external IDs with the same displayed label. That is unsafe: `ING-0014` and `ING-0069` must remain separate, and v22.5 candidate rows use `ING-0069` with the label `Шпик` while v22.13/PR39 calls the identity `Жир кулинарный`. DC1 records this as a semantic review blocker rather than guessing equivalence.

Existing PR39 FoodIngredient mappings are reused as identity decisions only. DC1 does **not** claim that a current nutrition profile is automatically suitable for every recipe source form; profile/form suitability remains explicitly not revalidated here.

## DECISION — proposed DC2 triage

{dc2_lines}

This is prioritization only. It does not authorize a production write or assert that authority/rights are closed.

## DECISION — proposed DC3 triage

{dc3_lines}

No source branch is selected merely because it reduces data debt. All batch membership is evidence triage only.

## Assortment review — preserve, do not auto-expand

Category counts in the original 68-family funnel:

{cat_lines}

The funnel is visibly skewed toward soups, vegetables/potatoes and egg dishes. Standalone meat/poultry breadth is limited. This is a DC3 variety blocker for a realistic family week, but DC1 does not expand the candidate universe automatically.

## Authority and rights boundary

- External v22.5/v22.13 nutrient values are not promoted to FamilyFoodOS production Nutrition.
- Existing PR39 mappings mean accepted identity mapping, not automatic nutrition-profile/form approval.
- New/form-split/proxy rows retain `NOT_VERIFIED_DC1`; `authority_search_strategy` is only the next verification strategy, not a verified source.
- Rights for future external authority rows remain `PENDING_SELECTED_SOURCE_REVIEW` until an exact source is selected and reviewed under the canonical DATA-CORPUS-V1 policy.

## Files

- `candidate-recipes.csv` — 68 candidates, full relationship counts, variant/choice/optional/boundary status and DC3 triage.
- `food-demand.csv` — one row per external ingredient identity; no unsafe name-based deduplication.
- `compatibility.csv` — v22.5 vs v22.13 count/semantic compatibility status.
- `source-relationships-part1.csv` … `part4.csv` — all {len(relationship_output)} retained source relationship/calculation rows for the 68 candidates, sharded by original v22.5 row file and excluding nutrient values.
- `batch-plan.json` — exact non-overlapping DC2/DC3 triage partitions.
- `source-artifacts.json` — input hashes, source usage limits and missing v22.13 row-shard blocker.
- `summary.json` — machine-readable reconciliation totals.
- `checksums.sha256` — deterministic hashes of generated package files.

## Rebuild

The package is generated from raw XLSX bytes plus the accepted PR #39 mapping package. Example:

```bash
python scripts/build_data_corpus_v1_dc1.py \
  --row-shard /path/russian_normative_recipes_v22_5_row_nutrients_part1.xlsx \
  --row-shard /path/russian_normative_recipes_v22_5_row_nutrients_part2.xlsx \
  --row-shard /path/russian_normative_recipes_v22_5_row_nutrients_part3.xlsx \
  --row-shard /path/russian_normative_recipes_v22_5_row_nutrients_part4.xlsx \
  --v22-5-manifest /path/russian_normative_recipes_v22_5_manifest.xlsx \
  --v22-13-mass /path/russian_normative_recipes_v22_13_mass_nutrients.xlsx \
  --v22-13-audit /path/russian_normative_recipes_v22_13_integrity_audit.xlsx \
  --v22-13-manifest /path/russian_normative_recipes_v22_13_manifest.xlsx \
  --mapping-dir data/curation/v22-13-map-a \
  --output /tmp/data-corpus-v1-dc1

python scripts/validate_data_corpus_v1_dc1.py /tmp/data-corpus-v1-dc1
```

Source XLSX files are not committed by this package; exact required SHA-256 values are enforced by the generator and recorded in `source-artifacts.json`.

## Verification invariants

The generator fails closed on:

- row-shard truncation or candidate loss;
- missing mapping IDs or duplicate mapping identities;
- duplicate source relationship rows;
- blank/non-positive candidate ingredient amounts;
- v22.5/v22.13 calculation-row or variant-block mismatch;
- disagreement with accepted PR39 per-recipe mapping aggregates;
- summary/batch partition omission or overlap;
- assignment of a simple/exact DC3 triage status while ChoiceGroup, optional, multi-variant or garnish/sauce boundary remains unresolved.

A second build from the same inputs must be byte-identical for all generated package files.

## OPEN blockers

1. **{len(review_candidates)}** candidate families still require variant/choice/optional/boundary review before exact publication-branch selection.
2. **{len(work_demands)}** external identities require DC2-level closure if their dependent recipes are pursued.
3. **{len(compat_gap)}** candidates have additional v22.13 non-calc relationship rows; exact v22.13 row shards were not available here, so full row-level compatibility remains unproven.
4. **{len(exact_label_diffs)}** used identities have a v22.5/v22.13 label difference requiring review; `ING-0069` is the clearest explicit conflict example.
5. The original 68-family funnel has assortment gaps for a realistic family week; DC1 records the gap but does not widen scope automatically.

## Stop condition

After review/merge of this evidence package, stop. DC2 and DC3 remain **NOT STARTED** and require separate reviewable production operations.
"""
    (out / "README.md").write_text(readme, encoding="utf-8")

    # All generated artifacts except checksum itself. Recovery directory is outside this output during generation.
    generated_names = [
        "README.md", "candidate-recipes.csv", "food-demand.csv", "compatibility.csv", *relationship_part_names,
        "batch-plan.json", "source-artifacts.json", "summary.json",
    ]
    checksum_lines = [f"{sha256(out / name)}  {name}" for name in generated_names]
    (out / "checksums.sha256").write_text("\n".join(checksum_lines) + "\n", encoding="utf-8")

    # Final self-reconciliation from serialized outputs, not in-memory assumptions.
    with (out / "candidate-recipes.csv").open(encoding="utf-8", newline="") as f:
        serialized_candidates = list(csv.DictReader(f))
    with (out / "food-demand.csv").open(encoding="utf-8", newline="") as f:
        serialized_demands = list(csv.DictReader(f))
    serialized_summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    serialized_plan = json.loads((out / "batch-plan.json").read_text(encoding="utf-8"))
    if len(serialized_candidates) != serialized_summary["candidate_selection"]["recipe_family_count"]:
        die("serialized summary candidate count mismatch")
    if len(serialized_demands) != serialized_summary["food_demand"]["demand_row_count"]:
        die("serialized summary demand count mismatch")
    for section, universe_field, id_field in [("dc2", serialized_demands, "external_ingredient_id"), ("dc3", serialized_candidates, "source_recipe_id")]:
        flattened = []
        key = "external_ingredient_ids" if section == "dc2" else "source_recipe_ids"
        for group in serialized_plan[section].values():
            flattened.extend(group[key])
        if sorted(flattened) != sorted(r[id_field] for r in universe_field):
            die(f"serialized {section} batch plan does not reconcile")

    print(json.dumps({
        "status": "PASS",
        "candidate_families": len(candidate_output),
        "source_relationship_rows": len(relevant_rows),
        "external_ingredient_ids": len(demand_rows),
        "accepted_existing_identity_mappings": len(existing_demands),
        "dc2_required_identities": len(work_demands),
        "one_source_variant": summary["candidate_selection"]["source_variant_structure"]["one_source_variant"],
        "multiple_source_variants": summary["candidate_selection"]["source_variant_structure"]["multiple_source_variants"],
        "structurally_simple": len(simple_candidates),
        "variant_or_boundary_review": len(review_candidates),
        "v22_13_extra_noncalc_relationship_gap_recipes": len(compat_gap),
        "semantic_label_review_ids": len(exact_label_diffs),
        "output": str(out),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
